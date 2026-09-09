from copy import copy
from io import BytesIO,StringIO
import pandas as pd
from openpyxl.styles import PatternFill
from fastapi import FastAPI,HTTPException,Response
from fastapi.middleware.cors import CORSMiddleware
from .engine import DEFAULT_RULES,EXPORT_COLUMNS,RuleEngine,export_rows
from .models import ComponentTemplate,ComponentTemplateInput,Project,ProjectInput,RuleConfig,ValidationResult
from .storage import Storage

app=FastAPI(title="Visual Wiring Definition API",version="2.0.0")
app.add_middleware(CORSMiddleware,allow_origins=["http://localhost:3000","http://localhost:3001","http://localhost:5173"],allow_origin_regex=r"https://.*\.chatgpt\.site",allow_methods=["*"],allow_headers=["*"])
storage=Storage()
def missing(kind,id):raise HTTPException(404,f"{kind} {id} not found")

@app.get("/health")
def health():return {"status":"ok","database":str(storage.path)}
@app.get("/api/rules",response_model=RuleConfig)
def rules():return DEFAULT_RULES

@app.get("/api/component-templates",response_model=list[ComponentTemplate])
def list_templates():return storage.list_templates()
@app.post("/api/component-templates",response_model=ComponentTemplate,status_code=201)
def create_template(value:ComponentTemplateInput):
    try:return storage.create_template(value)
    except Exception as error:
        if "UNIQUE" in str(error):raise HTTPException(409,f"Template {value.id} already exists")
        raise
@app.post("/api/component-templates/import",response_model=ComponentTemplate,status_code=201)
def import_template(value:ComponentTemplateInput):return create_template(value)
@app.get("/api/component-templates/{template_id}",response_model=ComponentTemplate)
def get_template(template_id:str):return storage.get_template(template_id) or missing("Template",template_id)
@app.put("/api/component-templates/{template_id}",response_model=ComponentTemplate)
def update_template(template_id:str,value:ComponentTemplateInput):return storage.update_template(template_id,value) or missing("Template",template_id)
@app.delete("/api/component-templates/{template_id}",status_code=204)
def delete_template(template_id:str):
    try:
        if not storage.delete_template(template_id):missing("Template",template_id)
    except ValueError as error:raise HTTPException(409,str(error))

@app.get("/api/projects",response_model=list[Project])
def list_projects():return storage.list_projects()
@app.post("/api/projects",response_model=Project,status_code=201)
def create_project(value:ProjectInput):return storage.create_project(value)
@app.get("/api/projects/{project_id}",response_model=Project)
def get_project(project_id:str):return storage.get_project(project_id) or missing("Project",project_id)
@app.put("/api/projects/{project_id}",response_model=Project)
def update_project(project_id:str,value:ProjectInput):return storage.update_project(project_id,value) or missing("Project",project_id)
@app.delete("/api/projects/{project_id}",status_code=204)
def delete_project(project_id:str):
    if not storage.delete_project(project_id):missing("Project",project_id)

def project_and_templates(project_id):return get_project(project_id),storage.list_templates()
@app.post("/api/projects/{project_id}/validate",response_model=ValidationResult)
def validate_project(project_id:str):
    project,templates=project_and_templates(project_id);return RuleEngine(project,templates).validate()
def checked_rows(project_id):
    project,templates=project_and_templates(project_id);result=RuleEngine(project,templates).validate()
    if not result.valid:raise HTTPException(422,detail=result.model_dump(by_alias=True))
    return export_rows(project,templates)
@app.get("/api/projects/{project_id}/export/csv")
def export_csv(project_id:str):
    out=StringIO();pd.DataFrame(checked_rows(project_id),columns=EXPORT_COLUMNS).to_csv(out,index=False)
    return Response(out.getvalue(),media_type="text/csv",headers={"Content-Disposition":f"attachment; filename={project_id}-wiring.csv"})
@app.get("/api/projects/{project_id}/export/excel")
def export_excel(project_id:str):
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        pd.DataFrame(checked_rows(project_id),columns=EXPORT_COLUMNS).to_excel(writer,sheet_name="Wiring Definition",index=False)
        ws=writer.book["Wiring Definition"];ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            font=copy(cell.font);font.bold=True;cell.font=font;cell.fill=PatternFill(fill_type="solid",fgColor="E9E3FF")
        for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=max(14,min(32,max(len(str(cell.value or "")) for cell in column)+2))
    return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f"attachment; filename={project_id}-wiring.xlsx"})
